"""Metadata-driven, vectorized asset comparison primitives.

This module deliberately separates metadata/formula/comparison logic from
source-specific extraction (PPC and pandapower net). Source adapters should
provide arrays keyed by parameter name or parameter ID.
"""
from __future__ import annotations

import ast
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Mapping

import numpy as np
import pandas as pd
from src.comparison.result_structures import MetricDefinition
from src.logger import logger


class FormulaError(ValueError):
    pass


class SafeFormula:
    """Restricted vectorized formula evaluator; never executes arbitrary Python."""
    FUNCTIONS = {
        "abs": np.abs,
        "sqrt": np.sqrt,
        "sum": np.sum,
        "mean": np.mean,
        "max": np.maximum,
        "min": np.minimum,
        "hypot": np.hypot,
    }

    def evaluate(self, expression: str, values: Mapping[str, Any]) -> Any:
        try:
            tree = ast.parse(expression, mode="eval")
        except SyntaxError as exc:
            raise FormulaError(f"Invalid formula {expression!r}: {exc.msg}") from exc
        return self._node(tree.body, values)

    def _node(self, node: ast.AST, values: Mapping[str, Any]) -> Any:
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.Name):
            if node.id not in values:
                raise FormulaError(f"Unknown formula reference: {node.id}")
            return values[node.id]
        if isinstance(node, ast.BinOp):
            left, right = self._node(node.left, values), self._node(node.right, values)
            ops = {ast.Add: np.add, ast.Sub: np.subtract, ast.Mult: np.multiply,
                   ast.Div: np.divide, ast.Pow: np.power}
            op = ops.get(type(node.op))
            if op is None:
                raise FormulaError(f"Unsupported operator: {type(node.op).__name__}")
            with np.errstate(divide="ignore", invalid="ignore", over="ignore"):
                return op(left, right)
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = self._node(node.operand, values)
            return value if isinstance(node.op, ast.UAdd) else np.negative(value)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            fn = self.FUNCTIONS.get(node.func.id)
            if fn is None or node.keywords:
                raise FormulaError(f"Unsupported function: {node.func.id}")
            args = [self._node(arg, values) for arg in node.args]
            return fn(*args)
        raise FormulaError(f"Unsupported formula expression: {type(node).__name__}")


def _delta(a, b, comparison):
    if comparison == "angle":
        return np.abs(((a - b + 180.0) % 360.0) - 180.0)
    if comparison == "absolute":
        return np.abs(a - b)
    raise ValueError("Exact comparisons do not have a numeric delta operation")


def _signed_delta(a, b, comparison):
    """
    Calculate signed differences using the convention Case B - Case A.

    Angular differences follow the shortest path around the
    360-degree boundary.
    """
    if comparison == "angle":
        return ((b - a + 180.0) % 360.0) - 180.0

    if comparison == "absolute":
        return b - a

    raise ValueError(
        "Exact comparisons do not have a numeric signed delta"
    )


def _distribution_stats(values, prefix):
    """Calculate distribution statistics for finite numeric values."""
    values = np.asarray(values, dtype=float).reshape(-1)
    values = values[np.isfinite(values)]

    metric_names = (
        "mean", "max", "min", "median",
        "p90", "p95", "p99", "mad",
    )

    if values.size == 0:
        return {f"{name}_{prefix}": None for name in metric_names}

    median = float(np.median(values))
    p90, p95, p99 = (
        float(value)
        for value in np.percentile(values, [90, 95, 99])
    )

    return {
        f"mean_{prefix}": float(np.mean(values)),
        f"max_{prefix}": float(np.max(values)),
        f"min_{prefix}": float(np.min(values)),
        f"median_{prefix}": median,
        f"p90_{prefix}": p90,
        f"p95_{prefix}": p95,
        f"p99_{prefix}": p99,
        f"mad_{prefix}": float(
            np.median(np.abs(values - median))
        ),
    }


def _relative_delta(
        delta: np.ndarray,
        a: np.ndarray,
        b: np.ndarray,
        valid: np.ndarray,
        mismatch: np.ndarray,
        tolerance: float,
) -> np.ndarray:
    """Calculate relative delta percentages using the existing comparison delta.

    ```
    Matching values are represented as 0%.
    A mismatching value with a near-zero denominator is undefined (NaN).
    """
    delta = np.asarray(delta, dtype=float)
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)

    denominator = np.maximum(np.abs(a), np.abs(b))
    result = np.full(delta.shape, np.nan, dtype=float)

    matching = valid & ~mismatch
    result[matching] = 0.0

    defined = valid & mismatch & (denominator > tolerance)
    result[defined] = (
        100.0 * delta[defined] / denominator[defined]
    )

    return result


def _asset_stats(asset_ids, mismatch, valid):
        """Calculate unique affected-asset statistics."""

        valid_ids = asset_ids[valid]
        affected_ids = asset_ids[valid & mismatch]

        # Exclude missing asset identifiers from the counts.
        valid_ids = valid_ids[~pd.isna(valid_ids)]
        affected_ids = affected_ids[~pd.isna(affected_ids)]

        asset_count = int(pd.Series(valid_ids).nunique())
        affected_count = int(pd.Series(affected_ids).nunique())

        return {
            "asset_count": asset_count,
            "affected_asset_count": affected_count,
            "affected_asset_percent": (
                100.0 * affected_count / asset_count
                if asset_count
                else None
            ),
        }


def _signed_delta_stats(signed_delta, mismatch, valid):
    """Calculate directional statistics for mismatching values."""

    signed_delta = np.asarray(signed_delta, dtype=float)
    mismatch = np.asarray(mismatch, dtype=bool)
    valid = np.asarray(valid, dtype=bool)

    signed = signed_delta[valid & mismatch]
    signed = signed[np.isfinite(signed)]

    if signed.size == 0:
        return {
            "mean_signed_delta": None,
            "median_signed_delta": None,
            "positive_delta_count": 0,
            "negative_delta_count": 0,
            "positive_delta_percent": None,
            "negative_delta_percent": None,
        }

    positive_count = int(np.count_nonzero(signed > 0))
    negative_count = int(np.count_nonzero(signed < 0))
    count = int(signed.size)

    return {
        "mean_signed_delta": float(np.mean(signed)),
        "median_signed_delta": float(np.median(signed)),
        "positive_delta_count": positive_count,
        "negative_delta_count": negative_count,
        "positive_delta_percent": 100.0 * positive_count / count,
        "negative_delta_percent": 100.0 * negative_count / count,
    }


def _stats(delta, mismatch, severe, valid):
    """Calculate comparison counts and distribution statistics."""
    d = np.asarray(delta)[valid]

    count = int(valid.sum())
    mismatches = int(mismatch.sum())
    severe_count = int(severe.sum())

    statistics = {
        "comparable_count": count,
        "mismatch_count": mismatches,
        "severe_mismatch_count": severe_count,
        "mismatch_percent": (
            100.0 * mismatches / count if count else None
        ),
    }

    statistics.update(_distribution_stats(d, "delta"))

    return statistics


def compare_arrays(definition, values_a, values_b, asset_ids = None):
    """Compare corresponding rows using vectorized NumPy operations."""
    a, b = np.asarray(values_a), np.asarray(values_b)
    if a.shape != b.shape:
        return {"scope": definition.scope, "asset_name": definition.asset_name,
                "parameter_name": definition.parameter_name, "display_name": definition.display_name,
                "unit": definition.unit, "status": "NOT PERFORMED",
                "reason": f"Shape mismatch: Case A {a.shape}, Case B {b.shape}",
                "case_a": a, "case_b": b, "delta": None, "statistics": {}}
    if definition.comparison == "exact":
        valid = pd.notna(a) & pd.notna(b)
        # Exact comparison should not round numeric values.
        equal = np.equal(a, b)
        mismatch = valid & ~equal
        delta = None
        severe = np.zeros(a.shape, dtype=bool)
        statistics = {"comparable_count": int(valid.sum()),
                      "mismatch_count": int(mismatch.sum()),
                      "severe_mismatch_count": 0,
                      "mismatch_percent": (100.0 * mismatch.sum() / valid.sum()) if valid.any() else None}
    else:
        try:
            af, bf = a.astype(float, copy=False), b.astype(float, copy=False)
        except (TypeError, ValueError):
            return {"scope": definition.scope, "asset_name": definition.asset_name,
                    "parameter_name": definition.parameter_name, "display_name": definition.display_name,
                    "unit": definition.unit, "status": "NOT PERFORMED",
                    "reason": "Values are not numeric", "case_a": a, "case_b": b,
                    "delta": None, "statistics": {}}
        valid = np.isfinite(af) & np.isfinite(bf)
        delta = _delta(af, bf, definition.comparison)
        mismatch = valid & (delta > definition.tolerance)
        severe = (valid & (delta > definition.severe_threshold)
                  if definition.severe_threshold is not None else np.zeros(a.shape, dtype=bool))
        statistics = _stats(delta, mismatch, severe, valid)
        relative_delta = _relative_delta(
            delta=delta,
            a=af,
            b=bf,
            valid=valid,
            mismatch=mismatch,
            tolerance=definition.tolerance,
        )

        statistics.update(_distribution_stats(relative_delta[valid], "relative_delta"))

        signed_delta = _signed_delta(a=af, b=bf, comparison=definition.comparison)

        statistics.update(_signed_delta_stats( signed_delta=signed_delta, mismatch=mismatch, valid=valid))

    statistics.update(_asset_stats(asset_ids=asset_ids, mismatch=mismatch, valid=valid))
    if not valid.any():
        status = "NOT PERFORMED"
        reason = "No valid corresponding values"
    elif severe.any():
        status, reason = "SEVERE MISMATCH", None
    elif mismatch.any():
        status, reason = "MISMATCH", None
    else:
        status, reason = "MATCH", None
    return {"scope": definition.scope, "asset_name": definition.asset_name,
            "asset_id": definition.asset_id, "parameter_name": definition.parameter_name,
            "parameter_id": definition.parameter_id, "display_name": definition.display_name,
            "unit": definition.unit, "comparison": definition.comparison,
            "tolerance": definition.tolerance, "severe_threshold": definition.severe_threshold,
            "asset_ids": asset_ids, "case_a": a, "case_b": b, "delta": delta,
            "mismatch_mask": mismatch, "severe_mask": severe, "status": status,
            "reason": reason, "statistics": statistics}


def compare_scope_asset(definitions, table_a, table_b, asset_ids):
    """
    Run once for a scope/asset group and compare all enabled parameters.
    """
    results = []
    for definition in definitions:
        try:
            if definition.source == "derived":
                a = SafeFormula().evaluate(definition.parameter_id, table_a)
                b = SafeFormula().evaluate(definition.parameter_id, table_b)
            else:
                a = table_a[definition.parameter_id]
                b = table_b[definition.parameter_id]
                if definition.asset_name == "load" and definition.scope == 'model':
                    a = a[np.flatnonzero((table_a['PD'] != 0) | (table_a['QD'] != 0))]
                    b = b[np.flatnonzero((table_b['PD'] != 0) | (table_b['QD'] != 0))]
                    asset_ids = np.asarray(list(range(len(a))))
            results.append(compare_arrays(definition, a, b, asset_ids=asset_ids))
        except (KeyError, AttributeError, IndexError, ValueError, TypeError) as exc:
            results.append({"scope": definition.scope, "asset_name": definition.asset_name,
                            "asset_id": definition.asset_id, "parameter_name": definition.parameter_name,
                            "parameter_id": definition.parameter_id, "display_name": definition.display_name,
                            "unit": definition.unit, "status": "NOT PERFORMED",
                            "reason": str(exc), "case_a": None, "case_b": None,
                            "delta": None, "statistics": {}})
    return results


def group_definitions(definitions):
    groups = {}
    for definition in definitions:
        groups.setdefault((definition.scope, definition.asset_name), []).append(definition)
    return groups


def export_comparison_excel(results, path):
    """Export long-form parameter/asset details and aggregate statistics.

    A row is emitted per asset and parameter; callers can provide array-valued
    results. The workbook has one sheet per scope/asset plus system_summary.
    """
    from openpyxl import Workbook
    from openpyxl.utils.dataframe import dataframe_to_rows

    workbook = Workbook(write_only=True)
    grouped: dict[str, list[dict[str, Any]]] = {}
    summaries = []
    for result in results:
        sheet = f"{result.get('scope','')}_{result.get('asset_name','asset')}"[:31]
        grouped.setdefault(sheet, []).append(result)
        summaries.append({"Scope": result.get("scope"), "Asset": result.get("asset_name"),
                          "Parameter": result.get("parameter_name"), "Status": result.get("status"),
                          "Reason": result.get("reason"), **result.get("statistics", {})})
    for sheet_name, entries in grouped.items():
        ws = workbook.create_sheet(sheet_name)
        rows = []
        for entry in entries:
            a, b, d = entry.get("case_a"), entry.get("case_b"), entry.get("delta")
            if a is None or b is None:
                rows.append({"Asset ID": None, "Parameter": entry.get("parameter_name"),
                             "Case A": None, "Case B": None, "Delta": None,
                             "Status": entry.get("status"), "Reason": entry.get("reason")})
                continue
            aa, bb = np.asarray(a).reshape(-1), np.asarray(b).reshape(-1)
            dd = np.asarray(d).reshape(-1) if d is not None else np.full(aa.shape, None, dtype=object)
            ids = entry.get("asset_ids")
            ids = np.arange(len(aa)) if ids is None else np.asarray(ids).reshape(-1)
            for i in range(len(aa)):
                rows.append({"Asset ID": ids[i] if i < len(ids) else i,
                             "Parameter": entry.get("parameter_name"), "Display Name": entry.get("display_name"),
                             "Unit": entry.get("unit"), "Case A": aa[i], "Case B": bb[i],
                             "Delta": dd[i] if i < len(dd) else None, "Status": entry.get("status")})
        frame = pd.DataFrame(rows)
        for row in dataframe_to_rows(frame, index=False, header=True):
            ws.append(row)
    ws = workbook.create_sheet("system_summary")
    frame = pd.DataFrame(summaries)
    for row in dataframe_to_rows(frame, index=False, header=True):
        ws.append(row)
    workbook.save(path)
